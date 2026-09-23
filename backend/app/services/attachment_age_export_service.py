"""
Generates the Attachment Age Excel export (GET /api/dashboard/attachment-age/export)
for the Overview page's Attachment Age card.

No Attachment Age calculation is reimplemented here. Every case-level row
comes from services.overview_period_metrics.get_overview_attachment_case_export_rows()
- an additive function that reuses the exact same _build_groups()/
_earliest_initiation_records()/_attachment_case_records() pieces already
powering the Overview Attachment Age card and its case-audit popover (see
that module for details). The workbook's Min/Avg/Max summary values are
computed FROM those same exported rows (mirroring
overview_period_metrics._attachment_stats()'s exact NaN/negative-minutes
exclusion rule for min/avg/max, while case counts include every row) - not
recalculated independently - so the export cannot diverge from what the
dashboard shows for the same period.
"""
from __future__ import annotations

import io
from typing import Any, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# Triggers the repo-root sys.path bootstrap as a side effect before the
# `services` package (which lives outside backend/) is imported below.
from app.core.config import settings  # noqa: F401

import services.overview_period_metrics as overview_period_metrics
from app.services.overview_service import _resolve_display_range, _resolve_range

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SECTION_FILL = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
_SECTION_FONT = Font(bold=True)
_THIN_BORDER = Border(*(Side(style="thin", color="D1D5DB") for _ in range(4)))
_DURATION_FORMAT = "[h]:mm:ss"


def _group_stats(rows: pd.DataFrame) -> tuple[int, float, float, float, float]:
    """Mirrors overview_period_metrics._attachment_stats()'s exact min/avg/
    max exclusion rule (NaN and negative minutes dropped from stats, but
    NOT from case_count), applied to an already-built case-level rows
    dataframe - so every summary number below (including Median, computed
    from this exact same filtered `minutes` series) is derived from the
    same rows the `data` sheet contains, not from a separately-called stat
    function. Returns (case_count, avg, min, max, median), all in minutes."""
    case_count = len(rows)
    if case_count == 0:
        return 0, 0.0, 0.0, 0.0, 0.0
    minutes = rows["_minutes"].dropna()
    minutes = minutes[minutes >= 0]
    if len(minutes) == 0:
        return case_count, 0.0, 0.0, 0.0, 0.0
    return (
        case_count,
        round(float(minutes.mean()), 1),
        round(float(minutes.min()), 1),
        round(float(minutes.max()), 1),
        round(float(minutes.median()), 1),
    )


def _minutes_to_day_fraction(minutes: float) -> float:
    """Excel stores time-of-day/duration values as a fraction of a day.
    Using the `[h]:mm:ss` number format (square brackets suppress the
    24-hour wraparound) lets a duration of many days display correctly as
    total hours:minutes:seconds, matching Attachment Age values that
    routinely exceed 24h."""
    return minutes / 1440.0


def _style_header_row(ws: Worksheet, row: int, last_col: int) -> None:
    for col in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _THIN_BORDER


def _write_stat_row(
    ws: Worksheet, row: int, label: str, delivery: str, stats: tuple[int, float, float, float, float]
) -> None:
    case_count, avg_minutes, min_minutes, max_minutes, median_minutes = stats
    ws.cell(row=row, column=1, value=label)
    ws.cell(row=row, column=2, value=delivery)
    ws.cell(row=row, column=3, value=case_count)
    # Column order: Case Count | Average | Min | Max | Median
    for col, minutes in ((4, avg_minutes), (5, min_minutes), (6, max_minutes), (7, median_minutes)):
        cell = ws.cell(row=row, column=col, value=_minutes_to_day_fraction(minutes))
        cell.number_format = _DURATION_FORMAT
    for col in range(1, 8):
        cell = ws.cell(row=row, column=col)
        cell.border = _THIN_BORDER
        if col in (3, 4, 5, 6, 7):
            cell.alignment = Alignment(horizontal="center")


def _write_section_header(ws: Worksheet, row: int, title: str, last_col: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=title)
    cell.fill = _SECTION_FILL
    cell.font = _SECTION_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    for col in range(1, last_col + 1):
        ws.cell(row=row, column=col).border = _THIN_BORDER


def _build_overview_sheet(
    wb: Workbook,
    inborn_nvd: pd.DataFrame,
    inborn_csection: pd.DataFrame,
    outborn_nvd: pd.DataFrame,
    outborn_csection: pd.DataFrame,
    all_rows: pd.DataFrame,
    period_label: str,
) -> None:
    ws = wb.create_sheet("Attachment age overview")
    last_col = 7

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title_cell = ws.cell(row=1, column=1, value="Attachment Age Overview")
    title_cell.font = Font(bold=True, size=14)

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    period_cell = ws.cell(row=2, column=1, value=period_label)
    period_cell.font = Font(italic=True, color="6B7280")

    header_row = 4
    headers = ["Group", "Delivery", "Case Count", "Average", "Min", "Max", "Median"]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)
    _style_header_row(ws, header_row, last_col)

    row = header_row + 1
    _write_section_header(ws, row, "Inborn", last_col)
    row += 1
    _write_stat_row(ws, row, "Inborn", "NVD", _group_stats(inborn_nvd))
    row += 1
    _write_stat_row(ws, row, "Inborn", "C-Section", _group_stats(inborn_csection))
    row += 1

    _write_section_header(ws, row, "Outborn", last_col)
    row += 1
    _write_stat_row(ws, row, "Outborn", "NVD", _group_stats(outborn_nvd))
    row += 1
    _write_stat_row(ws, row, "Outborn", "C-Section", _group_stats(outborn_csection))
    row += 1

    _write_section_header(ws, row, "Summary", last_col)
    row += 1
    inborn_rows = pd.concat([inborn_nvd, inborn_csection], ignore_index=True)
    outborn_rows = pd.concat([outborn_nvd, outborn_csection], ignore_index=True)
    _write_stat_row(ws, row, "Total", "NVD + C-Section", _group_stats(all_rows))
    row += 1
    _write_stat_row(ws, row, "Inborn", "NVD + C-Section", _group_stats(inborn_rows))
    row += 1
    _write_stat_row(ws, row, "Outborn", "NVD + C-Section", _group_stats(outborn_rows))
    row += 1
    _write_stat_row(ws, row, "Overall I+O (NVD + C-section)", "NVD + C-Section", _group_stats(all_rows))

    column_widths = [30, 16, 12, 12, 12, 12, 12]
    for col, width in enumerate(column_widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _build_data_sheet(wb: Workbook, rows: pd.DataFrame) -> None:
    """`Record ID` is the REDCap/database `recordid` of the DCT/daily-care
    form that supplied each case's earliest qualifying initiation timestamp
    (already present as `dct_recordid` on every row returned by
    overview_period_metrics.get_overview_attachment_case_export_rows() -
    see that function's _attachment_case_records() call - no new
    calculation). `Baby ID` (scr_babyid) is kept as its own column since it
    is still needed as the baby-level identifier."""
    ws = wb.create_sheet("data")

    headers = [
        "Record ID",
        "Baby ID",
        "Case (Inborn/Outborn)",
        "attachment age (hh:mm:ss)",
        "Hours",
        "Minutes",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=header)
    _style_header_row(ws, 1, len(headers))

    ordered = rows.sort_values(["cohort", "scr_babyid"], kind="stable").reset_index(drop=True)
    for i, record in enumerate(ordered.to_dict(orient="records")):
        row = i + 2
        minutes = record["_minutes"]
        ws.cell(row=row, column=1, value=record["dct_recordid"])
        ws.cell(row=row, column=2, value=record["scr_babyid"])
        ws.cell(row=row, column=3, value=record["cohort"])

        if minutes is None or pd.isna(minutes):
            ws.cell(row=row, column=4, value=None)
            ws.cell(row=row, column=5, value=None)
            ws.cell(row=row, column=6, value=None)
        else:
            minutes_val = float(minutes)
            if minutes_val >= 0:
                ws.cell(row=row, column=5, value=round(minutes_val / 60, 2))
                ws.cell(row=row, column=6, value=round(minutes_val, 1))
                duration_cell = ws.cell(row=row, column=4, value=_minutes_to_day_fraction(minutes_val))
                duration_cell.number_format = _DURATION_FORMAT
            else:
                # Physically impossible (initiation before birth) - always
                # caused by a baby with conflicting duplicate eligibility
                # records (different scr_dob/scr_tob per record; see the
                # same comment in services/indicators.py::_attachment_stats()
                # and services/overview_period_metrics.py's mirror). The
                # row is kept (same case_count as the dashboard, which also
                # counts it) but every duration-bearing cell is written as
                # flagged TEXT, never a plain number - so it cannot be
                # silently pulled into a SUM/AVERAGE run directly on this
                # sheet, while the real computed values stay visible for
                # audit. Not hidden, not zeroed, not recalculated from a
                # fabricated timestamp.
                ws.cell(row=row, column=4, value="Excluded (negative duration)")
                ws.cell(row=row, column=5, value=f"Excluded ({round(minutes_val / 60, 2)})")
                ws.cell(row=row, column=6, value=f"Excluded ({round(minutes_val, 1)})")

        for col in range(1, 7):
            ws.cell(row=row, column=col).border = _THIN_BORDER
            if col in (4, 5, 6):
                ws.cell(row=row, column=col).alignment = Alignment(horizontal="center")

    last_row = len(ordered) + 1
    ws.freeze_panes = "A2"
    if last_row >= 1:
        ws.auto_filter.ref = f"A1:F{last_row}"

    column_widths = [16, 16, 22, 24, 12, 12]
    for col, width in enumerate(column_widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def build_attachment_age_workbook(
    period: Optional[str] = None, from_date: Optional[str] = None, to_date: Optional[str] = None
) -> bytes:
    start, end = _resolve_range(period, from_date, to_date)
    display_start, display_end = _resolve_display_range(start, end)

    rows = overview_period_metrics.get_overview_attachment_case_export_rows(start, end)

    inborn_nvd = rows[(rows["cohort"] == "Inborn") & (rows["delivery"] == "NVD")]
    inborn_csection = rows[(rows["cohort"] == "Inborn") & (rows["delivery"] == "C-Section")]
    outborn_nvd = rows[(rows["cohort"] == "Outborn") & (rows["delivery"] == "NVD")]
    outborn_csection = rows[(rows["cohort"] == "Outborn") & (rows["delivery"] == "C-Section")]

    if display_start and display_end:
        period_label = f"Reporting Period: {display_start[:10]} to {display_end[:10]}"
    else:
        period_label = "Reporting Period: All Data"

    wb = Workbook()
    # Workbook() creates one default sheet - remove it so only the two
    # named sheets below exist, in the required order.
    wb.remove(wb.active)
    _build_overview_sheet(wb, inborn_nvd, inborn_csection, outborn_nvd, outborn_csection, rows, period_label)
    _build_data_sheet(wb, rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
